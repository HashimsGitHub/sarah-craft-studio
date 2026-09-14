#!/usr/bin/env python3
"""Import Sarah Craft Studio discount codes from CSV or Excel into MongoDB.

The import is idempotent: records are upserted by uppercase discount code.
PICKUPYYC is reserved for local pickup and is intentionally skipped.

Required environment variables:
  MONGODB_URI
Optional:
  MONGODB_DB (defaults to sarahcraftstudio)

Examples:
  python scripts/import_discounts.py discounts.xlsx --dry-run
  python scripts/import_discounts.py discounts.xlsx
  python scripts/import_discounts.py discounts.csv
"""

from __future__ import annotations

import argparse
import csv
import os
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from pymongo import MongoClient, UpdateOne

RESERVED_CODES = {"PICKUPYYC"}
ALLOWED_TYPES = {"percentage", "fixed", "free_shipping"}

ALIASES = {
    "code": {"code", "discount code", "coupon", "coupon code", "promo code", "promotion code"},
    "type": {"type", "discount type", "kind"},
    "value": {"value", "discount value", "amount", "percent", "percentage", "discount"},
    "minimumOrder": {"minimum order", "min order", "minimum purchase", "minimumorder"},
    "usageLimit": {"usage limit", "limit", "max uses", "maximum uses", "usagelimit"},
    "startsAt": {"starts", "start", "start date", "starts at", "startsat"},
    "expiresAt": {"expires", "expiry", "expiry date", "expiration date", "expires at", "expiresat"},
    "active": {"active", "enabled", "status"},
}


def norm_header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").replace("-", " ").split())


def canonical_header(value: Any) -> str | None:
    n = norm_header(value)
    for target, aliases in ALIASES.items():
        if n in aliases:
            return target
    return None


def parse_bool(value: Any, default: bool = True) -> bool:
    if value is None or str(value).strip() == "":
        return default
    return str(value).strip().lower() not in {"0", "false", "no", "n", "inactive", "disabled", "off"}


def parse_number(value: Any, default: float = 0) -> float:
    if value is None or str(value).strip() == "":
        return default
    text = str(value).strip().replace("$", "").replace("%", "").replace(",", "")
    return float(text)


def parse_optional_int(value: Any) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    return max(0, int(float(str(value).strip())))


def parse_date(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    raise ValueError(f"Unsupported date format: {text}")


def normalize_type(value: Any, raw_value: Any) -> str:
    text = str(value or "").strip().lower().replace("-", "_").replace(" ", "_")
    mapping = {
        "percent": "percentage",
        "percentage": "percentage",
        "%": "percentage",
        "fixed": "fixed",
        "fixed_amount": "fixed",
        "amount": "fixed",
        "dollar": "fixed",
        "free_shipping": "free_shipping",
        "shipping": "free_shipping",
    }
    if text in mapping:
        return mapping[text]
    # If the spreadsheet omits type but the value contains %, infer percentage.
    if not text and "%" in str(raw_value or ""):
        return "percentage"
    return "percentage" if not text else text


def read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def read_xlsx(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel import requires openpyxl: pip install openpyxl") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(x or "") for x in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:] if any(v not in (None, "") for v in row)]


def read_rows(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".csv":
        return read_csv(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    raise RuntimeError("Supported files are .csv, .xlsx and .xlsm")


def canonicalize_row(row: dict[str, Any]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for key, value in row.items():
        target = canonical_header(key)
        if target:
            out[target] = value
    return out


def build_discount(row: dict[str, Any]) -> dict[str, Any]:
    x = canonicalize_row(row)
    code = str(x.get("code") or "").strip().upper()
    if not code:
        raise ValueError("missing discount code")
    if code in RESERVED_CODES:
        raise ValueError("reserved pickup code; skipped")

    raw_value = x.get("value")
    dtype = normalize_type(x.get("type"), raw_value)
    if dtype not in ALLOWED_TYPES:
        raise ValueError(f"unsupported discount type: {dtype}")

    value = 0 if dtype == "free_shipping" else parse_number(raw_value, 0)
    if value < 0:
        raise ValueError("discount value cannot be negative")
    if dtype == "percentage" and value > 100:
        raise ValueError("percentage discount cannot exceed 100")

    now = datetime.utcnow()
    return {
        "code": code,
        "type": dtype,
        "value": value,
        "minimumOrder": max(0, parse_number(x.get("minimumOrder"), 0)),
        "usageLimit": parse_optional_int(x.get("usageLimit")),
        "startsAt": parse_date(x.get("startsAt")),
        "expiresAt": parse_date(x.get("expiresAt")),
        "active": parse_bool(x.get("active"), True),
        "updatedAt": now,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Import discount codes into Sarah Craft Studio MongoDB")
    parser.add_argument("file", type=Path, help="CSV or Excel file")
    parser.add_argument("--dry-run", action="store_true", help="Validate and preview without writing MongoDB")
    args = parser.parse_args()

    if not args.file.exists():
        print(f"ERROR: file not found: {args.file}", file=sys.stderr)
        return 2

    rows = read_rows(args.file)
    valid: list[dict[str, Any]] = []
    errors: list[str] = []
    seen: set[str] = set()

    for idx, row in enumerate(rows, start=2):
        try:
            d = build_discount(row)
            if d["code"] in seen:
                errors.append(f"row {idx}: duplicate code {d['code']} in source file")
                continue
            seen.add(d["code"])
            valid.append(d)
        except Exception as exc:
            errors.append(f"row {idx}: {exc}")

    print(f"Rows read: {len(rows)}")
    print(f"Valid discounts: {len(valid)}")
    print(f"Skipped/invalid: {len(errors)}")
    for msg in errors[:25]:
        print(f"  - {msg}")
    if len(errors) > 25:
        print(f"  ... and {len(errors) - 25} more")

    if args.dry_run:
        print("DRY RUN: no database changes made")
        for d in valid[:10]:
            print(f"  {d['code']}: {d['type']} {d['value']} active={d['active']}")
        return 0

    uri = os.environ.get("MONGODB_URI")
    if not uri:
        print("ERROR: set MONGODB_URI before running the import", file=sys.stderr)
        return 2

    client = MongoClient(uri, serverSelectionTimeoutMS=10000)
    db = client[os.environ.get("MONGODB_DB", "sarahcraftstudio")]
    collection = db["discounts"]

    operations: list[UpdateOne] = []
    for d in valid:
        update = dict(d)
        created = datetime.utcnow()
        operations.append(
            UpdateOne(
                {"code": d["code"]},
                {
                    "$set": update,
                    "$setOnInsert": {"createdAt": created, "usageCount": 0},
                },
                upsert=True,
            )
        )

    if operations:
        result = collection.bulk_write(operations, ordered=False)
        print(f"Inserted: {result.upserted_count}")
        print(f"Matched existing: {result.matched_count}")
        print(f"Modified existing: {result.modified_count}")
    else:
        print("No valid discounts to import")

    client.close()
    print("Import complete")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
